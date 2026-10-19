# SPDX-License-Identifier: AGPL-3.0-or-later
# Copyright (C) 2026 NORMATYS
"""Fiche de synthèse d'un calcul (cahier des charges, § 6.5) : document Word au papier de l'éditeur de la méthode,
converti en PDF par LibreOffice quand il est installé.

Mise en forme : Calibri 10 points, langue française dans les styles et les passages, titres portés par les styles
Titre, Titre 1 et Titre 2 aux couleurs de NORMATYS, pagination « Page X sur Y », tableaux à colonnes fixes, cellules
centrées verticalement, nombres centrés."""
from __future__ import annotations

import os
import shutil
import subprocess
import tempfile

STATIQUE = os.path.join(os.path.dirname(os.path.abspath(__file__)), 'statique')
BLEU, VERT, ENCRE, VOILE = '1A4F9C', '1E7A4A', '202335', 'EEF3F6'
MENTION = ('Méthode, conventions et seuils proposés, en attente de validation du Biodiversity Standard Council : '
           'les résultats ne sont pas opposables.')
NIVEAUX = {0: 'perte nette', 1: 'absence de perte nette', 2: 'gain net'}


class ErreurFiche(RuntimeError):
    """Fiche impossible à produire (bibliothèque ou LibreOffice absents)."""


def nombre(x, d=1) -> str:
    if x is None:
        return 'non défini'
    s = f'{x:,.{d}f}'.replace(',', ' ').replace('.', ',')
    return s.replace('-', '-')


def soffice() -> str | None:
    for c in (os.environ.get('OPENNAT_SOFFICE'), shutil.which('soffice'),
              r'C:\Program Files\LibreOffice\program\soffice.exe', '/usr/bin/soffice',
              '/Applications/LibreOffice.app/Contents/MacOS/soffice'):
        if c and os.path.exists(c):
            return c
    return None


def _langue(element, langue='fr-FR'):
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    rpr = element.get_or_add_rPr()
    lang = rpr.find(qn('w:lang'))
    if lang is None:
        lang = OxmlElement('w:lang')
        rpr.append(lang)
    for a in ('w:val', 'w:eastAsia', 'w:bidi'):
        lang.set(qn(a), langue)


def _champ(par, code):
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    r = par.add_run()
    for typ, txt in (('begin', None), (None, code), ('separate', None), (None, None), ('end', None)):
        if typ:
            e = OxmlElement('w:fldChar')
            e.set(qn('w:fldCharType'), typ)
            r._r.append(e)
        elif txt:
            e = OxmlElement('w:instrText')
            e.set(qn('xml:space'), 'preserve')
            e.text = txt
            r._r.append(e)
    return r


def _ombrage(cellule, couleur):
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    tc = cellule._tc.get_or_add_tcPr()
    sh = OxmlElement('w:shd')
    sh.set(qn('w:val'), 'clear')
    sh.set(qn('w:color'), 'auto')
    sh.set(qn('w:fill'), couleur)
    tc.append(sh)


def _tableau(doc, entetes, lignes, largeurs, numeriques=()):
    from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Cm, Pt, RGBColor
    t = doc.add_table(rows=1 + len(lignes), cols=len(entetes))
    t.style = 'Table Grid'
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    for i, ligne in enumerate([entetes] + lignes):
        for j, val in enumerate(ligne):
            c = t.cell(i, j)
            c.width = Cm(largeurs[j])
            c.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            p = c.paragraphs[0]
            p.paragraph_format.space_before = Pt(3)
            p.paragraph_format.space_after = Pt(3)
            run = p.add_run(str(val))
            if i == 0:
                run.bold = True
                run.font.color.rgb = RGBColor.from_string('FFFFFF')
                _ombrage(c, ENCRE)
            if i == 0 and j in numeriques or (i > 0 and j in numeriques):
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.keep_with_next = i < len(lignes)          # le tableau reste d'un bloc
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    for i, row in enumerate(t.rows):
        trpr = row._tr.get_or_add_trPr()
        e = OxmlElement('w:cantSplit')
        trpr.append(e)
        if i == 0:
            h = OxmlElement('w:tblHeader')
            trpr.append(h)
    return t


def construire_docx(res: dict, chemin: str):
    try:
        from docx import Document
        from docx.enum.text import WD_ALIGN_PARAGRAPH
        from docx.shared import Cm, Pt, RGBColor
    except ImportError as x:
        raise ErreurFiche('la bibliothèque python-docx est nécessaire à la fiche de synthèse') from x
    doc = Document()
    sec = doc.sections[0]
    sec.page_height, sec.page_width = Cm(29.7), Cm(21.0)
    sec.left_margin = sec.right_margin = Cm(2.0)
    sec.top_margin, sec.bottom_margin = Cm(2.2), Cm(2.0)
    st = doc.styles['Normal']
    st.font.name, st.font.size = 'Calibri', Pt(10)
    st.paragraph_format.space_before = st.paragraph_format.space_after = Cm(0.2)
    st.paragraph_format.line_spacing = 1.0
    _langue(st.element)
    for nom, taille, coul in (('Title', 16, BLEU), ('Heading 1', 13, BLEU), ('Heading 2', 11, VERT)):
        s = doc.styles[nom]
        s.font.name, s.font.size, s.font.bold = 'Calibri', Pt(taille), True
        s.font.color.rgb = RGBColor.from_string(coul)
        s.paragraph_format.space_before, s.paragraph_format.space_after = Cm(0.4), Cm(0.2)
        s.paragraph_format.keep_with_next = True
        _langue(s.element)
    # en-tête : logo de l'éditeur ; pied : pagination
    en = sec.header.paragraphs[0]
    en.alignment = WD_ALIGN_PARAGRAPH.LEFT
    en.add_run().add_picture(os.path.join(STATIQUE, 'normatys.png'), width=Cm(4.2))
    pied = sec.footer.paragraphs[0]
    pied.alignment = WD_ALIGN_PARAGRAPH.CENTER
    pied.add_run('OpenNAT · Fiche de synthèse · Page ')
    _champ(pied, 'PAGE')
    pied.add_run(' sur ')
    _champ(pied, 'NUMPAGES')
    for r in pied.runs:
        r.font.size = Pt(8)

    p, v = res['projet'], res['versions']
    t = doc.add_paragraph('Fiche de synthèse OpenNAT', style='Title')
    t.paragraph_format.space_before = Cm(1.0)
    doc.add_paragraph(f'{p.get("nom") or p["identifiant"]}')
    m = doc.add_paragraph()
    r = m.add_run(MENTION)
    r.italic = True
    r.font.color.rgb = RGBColor.from_string(VERT)

    doc.add_paragraph('Identification du calcul', style='Heading 1')
    _tableau(doc, ['Élément', 'Valeur'], [
        ['Projet', f'{p.get("nom") or ""} ({p["identifiant"]})'],
        ['Parcelle', f'{nombre(p["parcelle"], 0)} m²'],
        ['Typologie, zone biogéographique, zone climatique', f'{p["typologie"]}, {p["zone_biogeographique"]}, '
                                                             f'{p["zone_climatique"]}'],
        ['Méthode, moteur', f'{v["methode"]}, moteur {v["moteur"]}'],
        ['Version des données (SHA-256)', v['donnees']],
        ['Date du calcul, graine', f'{v["date"]}, {v["graine"]}'],
        ['Empreinte des entrées du RSNat', res['empreinte']['entrees']],
        ['Empreinte des sorties du RSNat', res['empreinte']['sorties']]], [5.5, 11.5])

    ref = next(x for x in res['variantes'] if x['reference'])
    s = ref['seuils']
    doc.add_paragraph('Attestation au dépôt du permis de construire', style='Heading 1')
    doc.add_paragraph(f'Variante de référence : {ref["nom"]}. L\'attestation ne porte que sur le plan masse (convention C-P1).')
    _tableau(doc, ['Indicateur', 'Projet', 'Seuil proposé', 'Respecté'], [
        ['Bnat (points)', nombre(ref['indicateurs']['bnat']), nombre(s['bnat_min']) if s['bnat_min'] is not None else 'sans seuil',
         {True: 'oui', False: 'non', None: 'sans objet'}[s['bnat_conforme']]],
        ['Pleine terre', f'{nombre(100 * ref["indicateurs"]["pleine_terre"])} %', f'{nombre(100 * s["pleine_terre_min"])} %',
         'oui' if s['pleine_terre_conforme'] else 'non'],
        ['Règles du PLU', '', '', {True: 'oui', False: 'non', None: 'sans objet'}[s['plu_conforme']]]],
        [5.0, 4.0, 4.0, 4.0], numeriques=(1, 2, 3))
    doc.add_paragraph('Conclusion : ' + ('le plan masse respecte les seuils proposés.' if s['permis'] else
                                         'le plan masse ne respecte pas tous les seuils ; l\'attestation ne conclut pas '
                                         'à la conformité.'))

    doc.add_paragraph('Indicateurs par variante', style='Heading 1')
    lignes = []
    for x in res['variantes']:
        i = x['indicateurs']
        lignes.append([x['nom'] + (' (référence)' if x['reference'] else ''), nombre(i['cnat']),
                       nombre(100 * i['cnat_relatif']) + ' %' if i['cnat_relatif'] is not None else 'non défini',
                       nombre(i['inat'], 0), nombre(i['snat'], 0), nombre(i['bnat']), NIVEAUX[i['niveau']]])
    _tableau(doc, ['Variante', 'Cnat (m²HE·an)', 'Cnat relatif', 'Inat (m²HE·an)', 'Snat (m²HE·an)', 'Bnat',
                   'Niveau du site'], lignes, [3.6, 2.3, 1.8, 2.5, 2.5, 1.5, 2.8], numeriques=(1, 2, 3, 4, 5, 6))
    doc.add_paragraph('Cnat : consommation de nature sur le site, intégrée sur 50 ans (positive : perte). Inat : '
                      'biodiversité grise des matériaux et du chantier. Snat = -(Cnat + Inat). Bnat : besoin de nature '
                      'de la conception, comparable au coefficient de biotope du PLU.')

    for x in res['variantes']:
        doc.add_paragraph(f'Variante {x["index"]} : {x["nom"]}', style='Heading 1')
        si, ma = x['site'], x['materiaux']
        doc.add_paragraph('Site', style='Heading 2')
        _tableau(doc, ['Grandeur', 'Valeur'], [
            ['Valeur du site sans projet (m²HE)', nombre(si['v0'])],
            ['Valeur intégrée sans projet, I_0 (m²HE·an)', nombre(si['i0'])],
            ['Valeur intégrée avec projet, I_1 (m²HE·an)', nombre(si['i1'])],
            ['Dette de maturité (m²HE·an)', nombre(si['dette'])],
            ['Destruction en période sensible (m²HE·an)', nombre(si['destruction'])],
            ['Arrosage annuel moyen (m³)', nombre(si['arrosage'])],
            ['Plan de gestion écologique', f'{nombre(si["plan"], 0)} ans']], [10.0, 7.0], numeriques=(1,))
        doc.add_paragraph('Biodiversité grise', style='Heading 2')
        _tableau(doc, ['Poste', 'Inat (m²HE·an)'], [
            ['Climat (émissions des composants et du chantier)', nombre(ma['inat_clim'], 0)],
            ['Bois (occupation des forêts)', nombre(ma['inat_bois'], 0)],
            ['Carrières (granulats)', nombre(ma['inat_carr'], 0)]], [10.0, 7.0], numeriques=(1,))
        inc = x.get('incertitude')
        if inc:
            doc.add_paragraph(f'Incertitude ({inc["tirages"]} tirages)', style='Heading 2')
            _tableau(doc, ['Indicateur', 'Centile 5', 'Médiane', 'Centile 95'],
                     [[i['indicateur'].replace('ecart_', 'écart à la référence, '), nombre(i['p5']), nombre(i['p50']),
                       nombre(i['p95'])] for i in inc['Intervalle']], [6.5, 3.5, 3.5, 3.5], numeriques=(1, 2, 3))
            doc.add_paragraph('Paramètres les plus influents sur Snat : ' + ' ; '.join(
                f'{q["cle"]} {q["parametre"]} ({q["table"]}, corrélation de rang {nombre(q["correlation"], 2)})'
                for q in inc['Parametre_Influent']) + '.')
    for par in doc.paragraphs:
        for run in par.runs:
            _langue(run._r)
    doc.core_properties.language = 'fr-FR'
    doc.core_properties.title = f'Fiche de synthèse OpenNAT, {p["identifiant"]}'
    doc.core_properties.author = 'Utilitaire OpenNAT (NORMATYS)'
    doc.save(chemin)


def produire(res: dict, dossier: str, ext: str) -> str:
    """Écrit la fiche dans le dossier de résultats ; ext : docx ou pdf."""
    nom = f'fiche_synthese_{res["projet"]["identifiant"]}'
    docx_ = os.path.join(dossier, nom + '.docx')
    construire_docx(res, docx_)
    if ext == 'docx':
        return docx_
    exe = soffice()
    if not exe:
        raise ErreurFiche('LibreOffice introuvable : la fiche est disponible en Word (.docx)')
    profil = 'file:///' + os.path.join(tempfile.gettempdir(), 'opennat-lo-profil').replace('\\', '/')
    subprocess.run([exe, f'-env:UserInstallation={profil}', '--headless', '--convert-to', 'pdf', '--outdir', dossier,
                    docx_], capture_output=True, timeout=180)
    pdf = os.path.join(dossier, nom + '.pdf')
    if not os.path.exists(pdf):
        raise ErreurFiche('la conversion en PDF par LibreOffice a échoué')
    return pdf
