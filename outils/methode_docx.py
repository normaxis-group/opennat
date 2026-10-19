# SPDX-License-Identifier: AGPL-3.0-or-later
# Copyright (C) 2026 NORMATYS
"""Produit la version Word du texte de la méthode (pour le Biodiversity Standard Council) et des documents
d'accompagnement (guides, journal des versions, notes).

    python outils/methode_docx.py methode/OpenNAT_v1.md <sortie.docx> [--papier normatys|irice-hors-portee]

Convertit le Markdown du dépôt (titres #, ##, ###, paragraphes, listes « - », tableaux, gras **)
puis applique le gabarit de Cédric PLANTAZ (Calibri 10, styles de titres, français, pagination)
et le papier NORMATYS. Outils : C:\\Users\\Cedric\\.claude\\outils\\gabarit_docx.py.
"""
import re
import sys

from docx import Document
from docx.shared import Cm

sys.path.insert(0, r'C:\Users\Cedric\.claude\outils')
from gabarit_docx import aerer, appliquer, controler, papier, remplissage, rendre_pdf, sauts_de_page  # noqa: E402

TABLEAU_INSECABLE = (12, 1200)   # lignes et caractères au plus pour qu'un tableau reste d'un bloc
MOTIF = re.compile(r'(\*\*[^*]+\*\*|`[^`]+`)')


def ecrire(par, texte):
    for morceau in MOTIF.split(texte):
        if not morceau:
            continue
        if morceau.startswith('**'):
            par.add_run(morceau[2:-2]).bold = True
        elif morceau.startswith('`'):
            par.add_run(morceau[1:-1])
        else:
            par.add_run(morceau)


def cellules(ligne):
    return [c.strip() for c in ligne.strip().strip('|').split('|')]


def largeurs(table, total_cm=17.0):
    """Largeurs fixes proportionnelles au contenu (le gabarit les égalise) : racine de la plus
    longue chaîne de chaque colonne, bornée, puis ramenée à la largeur utile."""
    from docx.oxml.ns import qn
    n = len(table.columns)
    poids = []
    for c in range(n):
        longueur = max(len(table.cell(r, c).text) for r in range(len(table.rows)))
        poids.append(min(7.0, max(1.6, longueur ** 0.5)))
    somme = sum(poids)
    cm = [total_cm * p / somme for p in poids]
    grille = table._tbl.tblGrid
    for c, col in enumerate(grille.findall(qn('w:gridCol'))):
        col.set(qn('w:w'), str(int(cm[c] / 2.54 * 1440)))
    for r in table.rows:
        for c, cel in enumerate(r.cells):
            cel.width = Cm(cm[c])
    tblpr = table._tbl.tblPr
    lay = tblpr.find(qn('w:tblLayout'))
    if lay is None:
        from docx.oxml import OxmlElement
        lay = OxmlElement('w:tblLayout')
        tblpr.append(lay)
    lay.set(qn('w:type'), 'fixed')
    # la ligne d'en-tête se répète sur chaque page où le tableau se poursuit
    from docx.oxml import OxmlElement
    trpr = table.rows[0]._tr.get_or_add_trPr()
    if trpr.find(qn('w:tblHeader')) is None:
        trpr.append(OxmlElement('w:tblHeader'))
    for r in table.rows:  # une ligne ne se coupe pas entre deux pages
        pr = r._tr.get_or_add_trPr()
        if pr.find(qn('w:cantSplit')) is None:
            pr.append(OxmlElement('w:cantSplit'))


def convertir(source, sortie, societe='normatys'):
    doc = Document()
    with open(source, encoding='utf-8') as f:
        lignes = f.read().splitlines()
    i = 0
    while i < len(lignes):
        l = lignes[i]
        if not l.strip() or l.strip() == '---':
            i += 1
            continue
        if l.strip().startswith('```'):
            i += 1
            while i < len(lignes) and not lignes[i].strip().startswith('```'):
                ecrire(doc.add_paragraph(), lignes[i].rstrip())
                i += 1
            i += 1
            continue
        if l.startswith('#'):
            niveau = len(l) - len(l.lstrip('#'))
            texte = l[niveau:].strip()
            style = {1: 'Title', 2: 'Heading 1', 3: 'Heading 2'}.get(niveau, 'Heading 3')
            par = doc.add_paragraph(style=style)
            ecrire(par, texte)
            if niveau == 2 and texte.startswith('Partie'):
                par.paragraph_format.page_break_before = True
            i += 1
            continue
        if l.lstrip().startswith('|'):
            bloc = []
            while i < len(lignes) and lignes[i].lstrip().startswith('|'):
                if not re.fullmatch(r'\s*\|[\s|:-]+\|\s*', lignes[i]):
                    bloc.append(cellules(lignes[i]))
                i += 1
            n = max(len(r) for r in bloc)
            t = doc.add_table(rows=len(bloc), cols=n)
            t.style = 'Table Grid'
            for r, ligne in enumerate(bloc):
                for c in range(n):
                    cel = t.cell(r, c)
                    ecrire(cel.paragraphs[0], ligne[c] if c < len(ligne) else '')
                    if r == 0:
                        for run in cel.paragraphs[0].runs:
                            run.bold = True
            continue
        if l.startswith('- '):
            ecrire(doc.add_paragraph(style='List Bullet'), l[2:].strip())
            i += 1
            continue
        # paragraphe : lignes consécutives non vides jusqu'à un bloc d'une autre nature
        texte = [l.strip()]
        i += 1
        while (i < len(lignes) and lignes[i].strip() and not lignes[i].startswith(('#', '- '))
               and not lignes[i].lstrip().startswith('|') and lignes[i].strip() != '---'):
            # une ligne d'équation numérotée ou d'énumération commence un nouveau paragraphe
            if re.match(r'\s*(\(\w+\)|\d+\. |[A-Z][\w ]{0,30} : )', lignes[i]):
                break
            texte.append(lignes[i].strip())
            i += 1
        ecrire(doc.add_paragraph(), ' '.join(texte))
    appliquer(doc, nature='note', detecter_titres=False)
    aerer(doc)
    papier(doc, societe)
    # un intertitre en gras (paragraphe entièrement gras et court) reste avec ce qui le suit
    for par in doc.paragraphs:
        if par.runs and len(par.text) < 120 and all(r.bold for r in par.runs if r.text.strip()):
            par.paragraph_format.keep_with_next = True
    # une phrase qui annonce une liste, un tableau ou une commande (« … : ») reste avec ce qui la suit
    for par in doc.paragraphs:
        if par.text.rstrip().endswith(':'):
            par.paragraph_format.keep_with_next = True
    for t in doc.tables:
        largeurs(t)
        volume = sum(len(c.text) for ligne in t.rows for c in ligne.cells)
        if len(t.rows) <= TABLEAU_INSECABLE[0] and volume <= TABLEAU_INSECABLE[1]:
            # un tableau court ne se coupe pas entre deux pages
            for ligne in t.rows[:-1]:
                for cel in ligne.cells:
                    for par in cel.paragraphs:
                        par.paragraph_format.keep_with_next = True
    titre = next((x[2:].strip() for x in lignes if x.startswith('# ')), '')
    if 'OpenNAT_v1' in source or not titre:
        titre = 'Méthode de calcul OpenNAT, version 1' + (', édition complète' if 'complete' in sortie else '')
    doc.core_properties.title = titre
    doc.core_properties.author = 'IRICE' if societe.startswith('irice') else 'NORMATYS'
    doc.core_properties.language = 'fr-FR'
    doc.save(sortie)
    return sortie


if __name__ == '__main__':
    societe = sys.argv[sys.argv.index('--papier') + 1] if '--papier' in sys.argv else 'normatys'
    s = convertir(sys.argv[1], sys.argv[2], societe)
    if '--sans-sauts' not in sys.argv:
        sauts_de_page(s)
    print(controler(s))
    pdf = rendre_pdf(s)
    print(pdf, remplissage(s))
