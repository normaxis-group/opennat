# SPDX-License-Identifier: AGPL-3.0-or-later
# Copyright (C) 2026 NORMATYS
"""Lettre de transmission de la méthode OpenNAT 1 au Biodiversity Standard Council et dossier de pièces.

    python outils/transmission_bsc.py

Produit livraison/transmission_bsc/ : la lettre (Word et PDF, papier NORMATYS, mise en page de courrier du gabarit),
le bordereau des pièces et une copie de chaque pièce en PDF. Rien n'est envoyé : la transmission se fait sur décision
de Cédric PLANTAZ, après le dépôt du nom OpenNAT.
"""
import os
import shutil
import sys

from docx import Document
from docx.shared import Cm

sys.path.insert(0, r'C:\Users\Cedric\.claude\outils')
from gabarit_docx import aerer, appliquer, controler, papier, rendre_pdf, remplissage  # noqa: E402

RACINE = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
LIV = os.path.join(RACINE, 'livraison')
SORTIE = os.path.join(LIV, 'transmission_bsc')
DATE = 'Aix-en-Provence, le 19 octobre 2026'
PIECES = [
    ('OpenNAT, dossier de soumission au Biodiversity Standard Council.pdf', 'Dossier de soumission : demande, vérifications, dix décisions, 70 conventions'),
    ('OpenNAT, pollution et biodiversité grise.pdf', 'Note de décision sur la pollution (décision D3)'),
    ('Méthode de calcul OpenNAT v1.pdf', 'Texte de la méthode OpenNAT 1'),
    ('Méthode de calcul OpenNAT v1, édition complète.pdf', 'Édition complète de la méthode (fiches algorithmes, données, fiches habitats, cas de validation)'),
    ("OpenNAT, contrôle du module G face à l'étude Carbone 4.pdf", "Contrôle d'ordre de grandeur de la biodiversité grise"),
    ('OpenNAT, guide utilisateur.pdf', "Guide utilisateur de l'outil de calcul"),
    ('OpenNAT, guide du contributeur.pdf', 'Guide du contributeur : proposer une valeur avec sa source'),
    ('OpenNAT, journal des versions.pdf', 'Journal des versions de la méthode, des données et du moteur'),
]


def pages(chemin):
    try:
        import fitz
        return len(fitz.open(chemin))
    except Exception:
        return None


def lettre():
    doc = Document()
    doc.add_paragraph('Mesdames et Messieurs les membres\ndu Biodiversity Standard Council\n'
                      'aux soins du secrétariat du comité, IRICE\n'
                      'Espace Wagner, Bâtiment A2\n10 rue du Lieutenant Parayre\n13290 Aix-en-Provence')
    doc.add_paragraph(DATE)
    doc.add_paragraph('Soumission de la méthode de calcul OpenNAT 1', style='Title')
    doc.add_paragraph('Mesdames, Messieurs,')
    doc.add_paragraph(
        "NORMATYS soumet au Biodiversity Standard Council la méthode de calcul OpenNAT 1. Elle compare sur cinquante ans la "
        "valeur écologique d'un site avant et après un projet de construction et y ajoute la nature consommée ailleurs par "
        "les matériaux et le chantier, en mètres carrés d'habitat équivalent. Comme pour nos référentiels, le comité valide "
        "la version définitive avant que NORMATYS l'adopte et la publie ; tant qu'il ne s'est pas prononcé, la méthode, ses "
        "seuils et ses résultats ne sont pas opposables.")
    doc.add_paragraph(
        "Le dossier de soumission joint dit ce que nous vous demandons : valider la méthode et ses 70 conventions, ou dire "
        "lesquelles reprendre ; trancher dix décisions, dont trois portent sur la biodiversité grise (bois, emprises "
        "industrielles, pollution) et changent les résultats ; fixer les seuils, aujourd'hui proposés. Chaque option est "
        "chiffrée par le moteur de calcul et accompagnée de notre recommandation ; une note distincte traite la pollution.")
    doc.add_paragraph(
        "Nous proposons deux séances. La première porterait sur la présentation de la méthode, ses vérifications et les "
        "décisions D4 et D6 à D9 ; la seconde, sur les décisions D1 à D3, D5 et D10 et sur la validation de la version. "
        "Le secrétariat du comité vous proposera des dates.")
    doc.add_paragraph(
        "Ces documents ne sont pas publics : la méthode, ses tables et son moteur ne seront diffusés qu'après leur "
        "adoption. Nous vous remercions de les réserver aux travaux du comité.")
    doc.add_paragraph('Nous vous prions de croire, Mesdames, Messieurs, à l\'assurance de notre considération distinguée.')
    sig = doc.add_paragraph('Pour NORMATYS\nNORMAXIS, présidente,\nreprésentée par Cédric PLANTAZ, gérant')
    sig.style = doc.styles['Normal']
    doc.add_paragraph('Pièces jointes', style='Heading 1')
    for nom, objet in PIECES:
        n = pages(os.path.join(LIV, nom))
        nombre = f'{n:,}'.replace(',', ' ') if n else ''
        doc.add_paragraph(f'{objet} ({nombre} pages)' if n else objet, style='List Bullet')
    appliquer(doc, nature='lettre', detecter_titres=False)
    sig.style = doc.styles['Signature']
    aerer(doc)
    for q in doc.paragraphs:                          # la formule de politesse reste avec la signature
        if q.text.startswith('Nous vous prions'):
            q.paragraph_format.keep_with_next = True
    papier(doc, 'normatys')
    doc.core_properties.title = 'Soumission de la méthode de calcul OpenNAT 1 au Biodiversity Standard Council'
    doc.core_properties.author = 'NORMATYS'
    doc.core_properties.language = 'fr-FR'
    chemin = os.path.join(SORTIE, '00_Lettre de transmission au Biodiversity Standard Council.docx')
    doc.save(chemin)
    print(controler(chemin))
    pdf = rendre_pdf(chemin)
    shutil.copyfile(pdf, chemin[:-5] + '.pdf')
    print(remplissage(chemin))
    return chemin


def pieces():
    for k, (nom, _) in enumerate(PIECES, 1):
        source = os.path.join(LIV, nom)
        if os.path.exists(source):
            shutil.copyfile(source, os.path.join(SORTIE, f'{k:02d}_{nom}'))
        else:
            print('pièce absente :', nom)


if __name__ == '__main__':
    os.makedirs(SORTIE, exist_ok=True)
    pieces()
    lettre()
    print(sorted(os.listdir(SORTIE)))
