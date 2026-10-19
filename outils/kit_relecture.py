# SPDX-License-Identifier: AGPL-3.0-or-later
# Copyright (C) 2026 NORMATYS
"""Classeur de relecture de la méthode OpenNAT par un écologue (condition de la décision D6 du dossier de soumission).

    python outils/kit_relecture.py

Écrit livraison/relecture_ecologue/OpenNAT, classeur de relecture écologue.xlsx : une feuille par objet à relire
(habitats, grilles d'état, conventions), les données de la version en vigueur à gauche, les colonnes du relecteur à
droite (avis, valeur ou libellé proposé, source, commentaire). Les retours entrent dans la méthode par le circuit du
guide du contributeur.
"""
import csv
import os
from collections import defaultdict

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

RACINE = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
DON = os.path.join(RACINE, 'donnees')
SORTIE = os.path.join(RACINE, 'livraison', 'relecture_ecologue')
FICHIER = os.path.join(SORTIE, 'OpenNAT, classeur de relecture écologue.xlsx')

ENCRE, VOILE, FILET, RELECTEUR = '202335', 'EEF3F6', 'C9D6DE', 'FFF4D6'
POLICE = 'Calibri'
AVIS = '"Conforme,À revoir,Sans avis"'
FAMILLES = ('habitats', 'sols', 'connectivité', 'temps', 'maturation et gestion', 'relevés', 'contexte, eau et seuils')


def lire(nom):
    with open(os.path.join(DON, nom), encoding='utf-8') as f:
        return list(csv.DictReader(f, delimiter=';'))


def nombre(x):
    try:
        v = float(x)
    except (TypeError, ValueError):
        return x
    return int(v) if v.is_integer() else round(v, 4)


def entete(ws, colonnes, relecteur_a_partir_de):
    trait = Side(style='thin', color=FILET)
    for j, (titre, largeur) in enumerate(colonnes, 1):
        c = ws.cell(row=1, column=j, value=titre)
        c.font = Font(name=POLICE, size=10, bold=True, color='FFFFFF')
        c.fill = PatternFill('solid', fgColor=ENCRE if j < relecteur_a_partir_de else '7A5B00')
        c.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
        c.border = Border(top=trait, bottom=trait, left=trait, right=trait)
        ws.column_dimensions[get_column_letter(j)].width = largeur
    ws.row_dimensions[1].height = 42
    ws.freeze_panes = ws.cell(row=2, column=3)
    ws.auto_filter.ref = f'A1:{get_column_letter(len(colonnes))}1'


def corps(ws, lignes, relecteur_a_partir_de, numeriques=()):
    trait = Side(style='thin', color=FILET)
    for i, ligne in enumerate(lignes, 2):
        for j, v in enumerate(ligne, 1):
            c = ws.cell(row=i, column=j, value=v)
            c.font = Font(name=POLICE, size=10)
            c.alignment = Alignment(vertical='center', wrap_text=True,
                                    horizontal='center' if j in numeriques else 'left')
            c.border = Border(top=trait, bottom=trait, left=trait, right=trait)
            if j >= relecteur_a_partir_de:
                c.fill = PatternFill('solid', fgColor=RELECTEUR)
    n = len(lignes) + 1
    dv = DataValidation(type='list', formula1=AVIS, allow_blank=True)
    dv.add(f'{get_column_letter(relecteur_a_partir_de)}2:{get_column_letter(relecteur_a_partir_de)}{n}')
    ws.add_data_validation(dv)
    ws.page_setup.orientation = 'landscape'
    ws.page_setup.fitToWidth, ws.page_setup.fitToHeight = 1, 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.print_title_rows = '1:1'
    ws.oddFooter.center.text = 'Page &P sur &N'
    ws.oddFooter.center.font = POLICE
    ws.oddFooter.center.size = 8
    ws.oddHeader.left.text = f'OpenNAT · relecture écologue · {ws.title}'
    ws.oddHeader.left.size = 8


def feuille_lisez_moi(wb):
    ws = wb.active
    ws.title = 'Lisez-moi'
    lignes = [
        ('Relecture écologique de la méthode OpenNAT 1', True, 14),
        ("Méthode éditée par NORMATYS, en attente de validation du Biodiversity Standard Council. Document de travail, non public.", False, 10),
        ('', False, 10),
        ('Ce qui est demandé', True, 11),
        ("Dire, pour chaque habitat, chaque critère de grille d'état et chaque convention, si la valeur ou le texte retenu vous paraît juste pour des opérations immobilières en France.", False, 10),
        ("Les colonnes surlignées sont les vôtres : Avis (Conforme, À revoir, Sans avis), valeur ou libellé proposé, source de votre proposition, commentaire.", False, 10),
        ("Une valeur proposée doit citer sa source : publication, référentiel, relevés. Sans source, l'avis reste utile mais ne peut pas entrer dans la méthode tel quel.", False, 10),
        ('', False, 10),
        ('Les feuilles', True, 11),
        ('Habitats : les 191 types de la nomenclature, leur valeur intrinsèque, leurs délais et leurs correspondances.', False, 10),
        ("Grilles d'état : les critères de chaque grille (critères dits essentiels signalés), les indicateurs notés de la grille des boisements et les règles de notation.", False, 10),
        ('Conventions : les conventions de méthode qui touchent aux habitats, aux sols, au temps et aux relevés.', False, 10),
        ('', False, 10),
        ('Repères', True, 11),
        ("Valeur intrinsèque : en m²HE par m², entre 0 et 1, issue de la valeur de biotope de l'ordonnance allemande de compensation (BKompV) ramenée à 1.", False, 10),
        ("Délais : années pour atteindre l'état bon, moyen ou mauvais après création (métrique statutaire anglaise).", False, 10),
        ("Difficulté de création : coefficient de la métrique anglaise (1 facile, 0,1 très difficile).", False, 10),
        ("Stades J, M, A : valeur des types ligneux jeunes, moyens et âgés.", False, 10),
        ('', False, 10),
        ('Retour', True, 11),
        ("Renvoyer le classeur complété à NORMATYS. Chaque proposition retenue sera inscrite au journal des versions avec sa source et votre nom, si vous l'acceptez.", False, 10),
    ]
    for i, (t, gras, taille) in enumerate(lignes, 1):
        c = ws.cell(row=i, column=1, value=t)
        c.font = Font(name=POLICE, size=taille, bold=gras, color=ENCRE)
        c.alignment = Alignment(wrap_text=True, vertical='top')
        lignes_texte = max(1, -(-len(t) // 88))                     # environ 88 caractères par ligne à cette largeur
        ws.row_dimensions[i].height = max(taille * 1.6, 14 * lignes_texte)
    ws.column_dimensions['A'].width = 92
    ws.page_setup.fitToWidth, ws.page_setup.fitToHeight = 1, 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.oddFooter.center.text = 'Page &P sur &N'
    ws.oddFooter.center.size = 8


def feuille_habitats(wb):
    nomen = lire('nomenclature_habitats.csv')
    hab = defaultdict(dict)
    for r in lire('habitats.csv'):
        hab[r['cle']][r['parametre']] = r
    ws = wb.create_sheet('Habitats')
    colonnes = [('Code OpenNAT', 10), ('Libellé', 34), ('EUNIS 2012', 9), ('Position', 13), ('Code BKompV', 11),
                ('Habitat de la métrique anglaise', 28), ("Grille d'état", 9), ('Valeur intrinsèque (m²HE/m²)', 11),
                ('Plage', 12), ('Stade J', 8), ('Stade M', 8), ('Stade A', 8), ('Délai état bon (an)', 9),
                ('Délai état moyen (an)', 9), ('Délai état mauvais (an)', 9), ('Difficulté de création', 10),
                ('Source de la valeur', 36), ('Justification de la correspondance', 46),
                ('Avis', 11), ('Valeur proposée', 11), ('Source proposée', 30), ('Commentaire', 40)]
    entete(ws, colonnes, 19)
    lignes = []
    for n in nomen:
        h = hab.get(n['code_vinat'], {})
        v = h.get('valeur_intrinseque', {})
        g = lambda p: nombre(h[p]['valeur']) if p in h else ''          # noqa: E731
        lignes.append([n['code_vinat'], n['libelle_fr'], n['code_eunis_2012'], n['position'], n['bkompv_code'],
                       n['uk_habitat'], n.get('grille_etat', ''), nombre(v.get('valeur', '')),
                       f"{nombre(v['incertitude_min'])} à {nombre(v['incertitude_max'])}" if v else '',
                       g('valeur_intrinseque_stade_J'), g('valeur_intrinseque_stade_M'), g('valeur_intrinseque_stade_A'),
                       g('delai_maturite_bon'), g('delai_maturite_moyen'), g('delai_maturite_mauvais'),
                       g('difficulte_creation'),
                       f"{v.get('source', '')} ; {v.get('reference_precise', '')}".strip(' ;'), n['justification'],
                       '', '', '', ''])
    corps(ws, lignes, 19, numeriques=(8, 9, 10, 11, 12, 13, 14, 15, 16))
    return len(lignes)


def feuille_grilles(wb):
    ws = wb.create_sheet("Grilles d'état")
    colonnes = [('Grille', 9), ('Critère ou paramètre', 16), ('Nature', 14), ('Libellé', 70), ('Source', 40),
                ('Avis', 11), ('Libellé ou valeur proposé', 50), ('Commentaire', 40)]
    entete(ws, colonnes, 6)
    lignes = []
    for r in lire('etat_criteres.csv'):
        par = r['parametre']
        if par.startswith('critere_'):
            m = par.split('_')
            code, nature = m[1], 'critère essentiel' if 'essentiel' in m else 'critère'
        elif par.startswith('indicateur_'):
            code, nature = par[len('indicateur_'):].replace('_', ' '), 'indicateur noté de 1 à 3'
        else:
            code, nature = f"{par.replace('_', ' ')} = {nombre(r['valeur'])}", 'règle de notation'
        lignes.append([r['cle'], code, nature, r['justification'], f"{r['source']} ; {r['reference_precise']}",
                       '', '', ''])
    corps(ws, lignes, 6)
    return len(lignes)


def feuille_conventions(wb):
    ws = wb.create_sheet('Conventions')
    colonnes = [('N°', 8), ('Famille', 18), ('Catégorie', 18), ('Lignes de données', 10), ('Énoncé', 90),
                ('Avis', 11), ('Rédaction proposée', 50), ('Commentaire', 40)]
    entete(ws, colonnes, 6)
    lignes = [[r['numero'], r['famille'], r['categorie'], nombre(r['lignes_de_donnees']), r['enonce'], '', '', '']
              for r in lire('conventions.csv') if r['famille'] in FAMILLES]
    corps(ws, lignes, 6, numeriques=(4,))
    return len(lignes)


def main():
    os.makedirs(SORTIE, exist_ok=True)
    wb = Workbook()
    feuille_lisez_moi(wb)
    n_h, n_g, n_c = feuille_habitats(wb), feuille_grilles(wb), feuille_conventions(wb)
    p = wb.properties
    p.title, p.creator, p.language = 'OpenNAT, classeur de relecture écologue', 'NORMATYS', 'fr-FR'
    p.subject = 'Relecture des habitats, des grilles d\'état et des conventions de la méthode OpenNAT 1'
    wb.save(FICHIER)
    print(FICHIER, f'{n_h} habitats, {n_g} critères, {n_c} conventions')


if __name__ == '__main__':
    main()
